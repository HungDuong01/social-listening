"""Dependency-free smoke test for the parts that do not need a browser.

Run with ``python scripts/selftest.py``. It exercises number parsing, metric
merge semantics, idempotent persistence, comment parent resolution and the
exporters against a throwaway database.
"""

from __future__ import annotations

import asyncio
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

_TEMP_DIR = tempfile.mkdtemp(prefix="social-crawler-selftest-")
# Pin the database explicitly. Overriding only SC_DATA_DIR is not enough: an
# SC_DATABASE_URL in .env wins over it, so once the project pointed at
# PostgreSQL this suite happily created its fixtures in the real database.
os.environ["SC_DATABASE_URL"] = (
    f"sqlite+aiosqlite:///{Path(_TEMP_DIR).as_posix()}/selftest.sqlite3"
)
os.environ["SC_DATA_DIR"] = _TEMP_DIR
os.environ["SC_EXPORTS_DIR"] = _TEMP_DIR
os.environ["SC_SESSIONS_DIR"] = _TEMP_DIR
os.environ["SC_LOG_LEVEL"] = "WARNING"
# Keep the inter-post pause out of the test runtime.
os.environ["SC_POST_DELAY_MIN_MS"] = "1"
os.environ["SC_POST_DELAY_MAX_MS"] = "2"

from app.core.enums import MetricSource, Platform, SearchFilter, TargetStatus  # noqa: E402
from app.core.logging import configure_logging  # noqa: E402
from app.core.utils import parse_compact_number  # noqa: E402
from app.db.repositories import (  # noqa: E402
    CrawlRepository,
    JobRepository,
    PostFilters,
    QueryRepository,
)
from app.db.session import dispose_db, init_db, session_scope  # noqa: E402
from app.domain.entities import (  # noqa: E402
    AuthorRef,
    CrawlResult,
    Metrics,
    RawComment,
    RawPost,
)
from app.services.export_service import ExportService  # noqa: E402

PASSED = 0
FAILED: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    global PASSED
    if condition:
        PASSED += 1
        print(f"  PASS  {label}")
    else:
        FAILED.append(label)
        print(f"  FAIL  {label} {detail}")


# ------------------------------------------------------------------- parsing


def test_number_parsing() -> None:
    print("\n[1] compact number parsing")
    cases = [
        ("1.9K", 1_900, True),
        ("2K", 2_000, True),
        ("342", 342, False),
        ("1,234", 1_234, False),
        ("1,9 N", 1_900, True),
        ("12.3M", 12_300_000, True),
        ("1.234.567", 1_234_567, False),
    ]
    for raw, expected, approximate in cases:
        parsed = parse_compact_number(raw)
        check(
            f"{raw!r} -> {expected}",
            parsed.value == expected and parsed.approximate == approximate,
            f"(got {parsed.value}, approx={parsed.approximate})",
        )

    check("empty string -> None", parse_compact_number("").value is None)
    check("'Translate' -> None", parse_compact_number("Translate").value is None)
    check("None -> None", parse_compact_number(None).value is None)
    check("int passthrough stays exact", parse_compact_number(0) == parse_compact_number(0))
    check("zero is a real value, not None", parse_compact_number(0).value == 0)


def test_metric_merge() -> None:
    print("\n[2] metric merge semantics")
    exact = Metrics(likes=1903, comments=31, source=MetricSource.API)
    rounded = Metrics(
        likes=1900, comments=31, reposts=50, shares=2000,
        source=MetricSource.DOM, approximate=True,
    )

    merged = exact.merged_with(rounded)
    check("exact API value survives the merge", merged.likes == 1903)
    check("gap is filled from the fallback", merged.reposts == 50)
    check("merge is marked mixed", merged.source is MetricSource.MIXED)
    check("merge inherits approximation flag", merged.approximate is True)

    zero = Metrics(reposts=0, source=MetricSource.API)
    filled = zero.merged_with(Metrics(reposts=50, source=MetricSource.DOM))
    check("a real 0 is not treated as missing", filled.reposts == 0)

    check("empty metrics detected", Metrics().is_empty is True)
    check("metrics with a 0 are not empty", Metrics(likes=0).is_empty is False)


# ---------------------------------------------------------------- persistence


def _sample_post(likes: int | None, comments: int | None) -> RawPost:
    return RawPost(
        platform=Platform.THREADS,
        external_id="ABC123",
        url="https://www.threads.com/@tester/post/ABC123",
        author=AuthorRef(platform=Platform.THREADS, username="tester", external_id="99"),
        content="Món mới của Phê La ngon quá",
        shortcode="ABC123",
        metrics=Metrics(
            likes=likes, comments=comments, reposts=50, shares=2000,
            source=MetricSource.API,
        ),
        hashtags=["phela"],
    )


def _sample_comments() -> list[RawComment]:
    author = AuthorRef(platform=Platform.THREADS, username="fan01")
    replier = AuthorRef(platform=Platform.THREADS, username="tester")
    return [
        # Deliberately out of order: the nested reply is persisted BEFORE the
        # comment it answers, which is what the platforms actually do.
        RawComment(
            platform=Platform.THREADS,
            external_id="C2",
            post_external_id="ABC123",
            author=replier,
            content="Cảm ơn bạn nhé",
            parent_external_id="C1",
            depth=2,
            is_author_reply=True,
            like_count=2,
        ),
        RawComment(
            platform=Platform.THREADS,
            external_id="C1",
            post_external_id="ABC123",
            author=author,
            content="Ngon thật đó",
            depth=1,
            like_count=9,
        ),
    ]


async def test_persistence() -> int:
    print("\n[3] persistence + idempotency")
    await init_db()

    async with session_scope() as session:
        job = await JobRepository(session).create_job(
            keyword="phê la",
            platforms=[Platform.THREADS],
            search_filter=SearchFilter.TOP,
            max_posts=10,
            max_comments_per_post=100,
            include_comments=True,
        )
        job_id = job.id

    async with session_scope() as session:
        repository = CrawlRepository(session)
        post_id, written = await repository.persist_result(
            job_id=job_id,
            result=CrawlResult(post=_sample_post(1903, 31), comments=_sample_comments(),
                               notes=["da lay du 2/2 comment"]),
            rank=1,
        )
    check("two comments stored", written == 2)

    # Re-crawl the same post with *less* information than the first pass.
    async with session_scope() as session:
        repository = CrawlRepository(session)
        second_id, _ = await repository.persist_result(
            job_id=job_id,
            result=CrawlResult(post=_sample_post(None, None), comments=_sample_comments()),
            rank=1,
        )
    check("re-crawl updates the same row", second_id == post_id)

    async with session_scope() as session:
        query = QueryRepository(session)
        posts, total = await query.list_posts(PostFilters(job_id=job_id))
        check("no duplicate post rows", total == 1, f"(total={total})")

        stored = posts[0]
        check("a None metric did not erase the stored value", stored.likes == 1903)
        check("non-null metrics still update", stored.reposts == 50)
        check("author linked", stored.author is not None and stored.author.username == "tester")
        check("comment counter refreshed", stored.comments_collected == 2)

        comments = await query.list_comments(stored.id)
        check("no duplicate comment rows", len(comments) == 2, f"(n={len(comments)})")

        by_id = {comment.external_id: comment for comment in comments}
        parent = by_id["C1"]
        child = by_id["C2"]
        check(
            "reply linked to its parent despite arriving first",
            child.parent_id == parent.id,
            f"(parent_id={child.parent_id}, expected {parent.id})",
        )
        check("author reply flagged", child.is_author_reply is True)

        stats = await query.stats()
        check("stats count posts", stats["posts"] == 1)
        check("stats count comments", stats["comments"] == 2)

    return job_id


async def test_job_status_rollup() -> None:
    print("\n[4] job status roll-up")
    from app.core.enums import JobStatus, TargetStatus

    async with session_scope() as session:
        repository = JobRepository(session)
        job = await repository.create_job(
            keyword="rollup",
            platforms=[Platform.THREADS, Platform.TIKTOK],
            search_filter=SearchFilter.TOP,
            max_posts=5,
            max_comments_per_post=5,
            include_comments=False,
        )
        await repository.finish_target(job.id, Platform.THREADS, TargetStatus.COMPLETED)
        await repository.finish_target(job.id, Platform.TIKTOK, TargetStatus.FAILED)
        status = await repository.derive_job_status(job.id)
    check("one ok + one failed -> partial", status is JobStatus.PARTIAL, f"(got {status})")

    async with session_scope() as session:
        repository = JobRepository(session)
        job = await repository.create_job(
            keyword="rollup2",
            platforms=[Platform.THREADS],
            search_filter=SearchFilter.TOP,
            max_posts=5,
            max_comments_per_post=5,
            include_comments=False,
        )
        await repository.finish_target(job.id, Platform.THREADS, TargetStatus.NEEDS_AUTH)
        status = await repository.derive_job_status(job.id)
    check("only needs_auth -> failed", status is JobStatus.FAILED, f"(got {status})")


async def test_exports(job_id: int) -> None:
    print("\n[5] exporters")
    service = ExportService()
    async with session_scope() as session:
        for fmt in ("xlsx", "csv", "json"):
            path = await service.export(
                session, filters=PostFilters(job_id=job_id), fmt=fmt, label="selftest"
            )
            check(f"{fmt} written", path.exists() and path.stat().st_size > 0, f"({path})")

    from openpyxl import load_workbook

    written = sorted(Path(_TEMP_DIR).glob("selftest_*.xlsx"))
    check("export lands under the keyword-led name", bool(written), "(no selftest_*.xlsx)")
    xlsx = written[-1]
    workbook = load_workbook(xlsx)
    check(
        "workbook matches the reference skeleton",
        set(workbook.sheetnames)
        == {"Report", "Report (Simple)", "Summary", "Comments", "Notes"},
        f"({workbook.sheetnames})",
    )
    check("report sheet has data rows", workbook["Report"].max_row >= 3)

    headers = [cell.value for cell in workbook["Report"][1]]
    check("Platform column present and second", headers[1] == "Platform", f"({headers[:3]})")
    for column in ("Note", "Comments Count (platform)", "Comments Collected"):
        check(f"Report has a {column!r} column", column in headers, f"({headers})")

    summary = workbook["Summary"]
    check("Summary is one row per post", summary.max_row == 2, f"(rows={summary.max_row})")
    check(
        "Summary carries Platform too",
        [c.value for c in summary[1]][1] == "Platform",
    )
    check("Notes sheet is populated", workbook["Notes"].max_row > 5)


def test_json_body_parsing() -> None:
    print("\n[6] tolerant JSON body parsing")
    from app.crawler.harvester import loads_multi

    check("plain object", loads_multi('{"a":1}') == [{"a": 1}])
    check(
        "Facebook anti-hijack prefix stripped",
        loads_multi('for (;;);{"a":1}') == [{"a": 1}],
    )
    check(
        "concatenated documents split",
        loads_multi('{"a":1}\n{"b":2}') == [{"a": 1}, {"b": 2}],
    )
    check("html body ignored", loads_multi("<!doctype html><html></html>") == [])
    check("empty body ignored", loads_multi("") == [])


def test_threads_extraction() -> None:
    """The two traps found while crawling Threads by hand."""
    print("\n[7] Threads extraction rules")
    from app.crawler.adapters.meta_common import is_media_node, node_identity
    from app.crawler.adapters.threads import ThreadsAdapter
    from app.crawler.miner import collect_nodes

    def post_node(code, username, text, likes, extra_thread_info=None, **kwargs):
        node = {
            "pk": f"pk_{code}",
            "code": code,
            "taken_at": 1_757_000_000,
            "like_count": likes,
            "user": {"pk": f"u_{username}", "username": username, "is_verified": False},
            "caption": {"text": text},
            "text_post_app_info": extra_thread_info or {},
        }
        node.update(kwargs)
        return node

    root = post_node(
        "ABC123", "tester", "Món mới của Phê La #phela", 1903,
        {"direct_reply_count": 31, "repost_count": 50, "reshare_count": 2000},
    )
    reply = post_node(
        "R1", "fan01", "Ngon thật đó", 9,
        {"reply_to_author": {"username": "tester"}, "direct_reply_count": 1},
    )
    nested = post_node(
        "R2", "tester", "Cảm ơn bạn nhé", 2,
        {"reply_to_author": {"username": "fan01"}},
    )
    # A "Related threads" entry: a standalone post with NO reply_to_author.
    related = post_node("REL1", "someoneelse", "Bài không liên quan", 500, {})
    # A reply belonging to a *different* conversation that Threads renders in
    # the trending feed on the same page. It does carry reply_to_author, so the
    # only thing separating it from a real reply is who it answers.
    foreign = post_node(
        "FOR1", "myqnam", "Bình luận của luồng khác", 2024,
        {"reply_to_author": {"username": "stranger"}},
    )

    payload = {
        "data": {
            "data": {
                "edges": [
                    {"node": {"thread_items": [{"post": root}]}},
                    {"node": {"thread_items": [{"post": reply}, {"post": nested}]}},
                    {"node": {"thread_items": [{"post": related}]}},
                    {"node": {"thread_items": [{"post": foreign}]}},
                ]
            }
        }
    }

    nodes = collect_nodes([payload], is_media_node, key=node_identity)
    check("all five post nodes mined from nested payload", len(nodes) == 5, f"(got {len(nodes)})")

    adapter = ThreadsAdapter()
    root_node = next(node for node in nodes if node_identity(node) == "ABC123")
    post = adapter._post_from_node(root_node, fallback_url="https://x/")

    check("exact likes from JSON (not rounded)", post.metrics.likes == 1903)
    check("direct_reply_count -> comments", post.metrics.comments == 31)
    check("repost_count -> reposts", post.metrics.reposts == 50)
    check("reshare_count -> shares", post.metrics.shares == 2000)
    check("metrics flagged as exact API values", post.metrics.source is MetricSource.API)
    check("metrics not marked approximate", post.metrics.approximate is False)
    check("post url rebuilt from author + code", post.url.endswith("/@tester/post/ABC123"))
    check("hashtag extracted", post.hashtags == ["phela"])

    comments, truncated = adapter._replies_from_nodes(
        nodes,
        root_node=root_node,
        root_external_id="ABC123",
        root_author="tester",
        limit=100,
    )
    ids = [comment.external_id for comment in comments]
    check("'Related threads' post excluded from comments", "REL1" not in ids, f"(ids={ids})")
    check(
        "reply from another conversation excluded",
        "FOR1" not in ids,
        f"(ids={ids})",
    )
    check("genuine replies kept", ids == ["R1", "R2"], f"(ids={ids})")
    check("not truncated", truncated is False)

    top, deep = comments[0], comments[1]
    check("top-level reply has depth 1", top.depth == 1)
    check("top-level reply has no comment parent", top.parent_external_id is None)
    check("nested reply has depth 2", deep.depth == 2)
    check("nested reply linked to the right parent", deep.parent_external_id == "R1")
    check("reply by the post author is flagged", deep.is_author_reply is True)
    check("reply by someone else is not flagged", top.is_author_reply is False)

    limited, truncated_flag = adapter._replies_from_nodes(
        nodes, root_node=root_node, root_external_id="ABC123", root_author="tester", limit=1
    )
    check("limit respected", len(limited) == 1 and truncated_flag is True)

    # 0 must mean "no limit". A plain `len(comments) >= limit` test is true on
    # the very first pass when limit is 0, which collects nothing at all.
    unlimited, unlimited_truncated = adapter._replies_from_nodes(
        nodes, root_node=root_node, root_external_id="ABC123", root_author="tester", limit=0
    )
    check(
        "limit=0 collects every reply, not zero",
        len(unlimited) == 2 and unlimited_truncated is False,
        f"(got {len(unlimited)}, truncated={unlimited_truncated})",
    )


def test_search_url_building() -> None:
    print("\n[8] search URL building")
    from app.crawler import get_adapter
    from app.crawler.base import CrawlOptions

    options = CrawlOptions(search_filter=SearchFilter.TOP)
    expectations = {
        Platform.THREADS: "threads.com/search?q=ph%C3%AA+la",
        Platform.INSTAGRAM: "instagram.com/explore/search/keyword/?q=ph%C3%AA+la",
        Platform.TIKTOK: "tiktok.com/search?q=ph%C3%AA+la",
        Platform.FACEBOOK: "facebook.com/search/posts/?q=ph%C3%AA+la",
    }
    for platform, fragment in expectations.items():
        url = get_adapter(platform).build_search_url("phê la", options)
        check(f"{platform.value} search url", fragment in url, f"(got {url})")

    tag_url = get_adapter(Platform.INSTAGRAM).build_search_url("#phela", options)
    check("instagram hashtag routes to /explore/tags/", "/explore/tags/phela/" in tag_url, tag_url)

    recent = get_adapter(Platform.THREADS).build_search_url(
        "abc", CrawlOptions(search_filter=SearchFilter.RECENT)
    )
    check("threads recent filter applied", "filter=recent" in recent, recent)


async def test_session_cookie_rules() -> None:
    """The exact bug that made Facebook logins loop forever."""
    print("\n[9] session cookie rules")
    from app.crawler import get_adapter

    class FakeContext:
        def __init__(self, names: dict[str, str]) -> None:
            self._cookies = [{"name": k, "value": v} for k, v in names.items()]

        async def cookies(self, _urls):  # noqa: ANN001 - mirrors Playwright's API
            return self._cookies

    facebook = get_adapter(Platform.FACEBOOK)

    only_id = FakeContext({"c_user": "123", "datr": "abc"})
    check(
        "c_user alone is NOT a session (it is set mid-challenge)",
        await facebook.has_live_session(only_id) is False,
    )

    full = FakeContext({"c_user": "123", "xs": "tok", "datr": "abc"})
    check("c_user + xs is a session", await facebook.has_live_session(full) is True)

    checkpointed = FakeContext({"c_user": "123", "xs": "tok", "checkpoint": "1"})
    check(
        "a checkpoint cookie blocks the session",
        await facebook.has_live_session(checkpointed) is False,
    )

    empty_value = FakeContext({"c_user": "123", "xs": ""})
    check(
        "an empty cookie value does not count",
        await facebook.has_live_session(empty_value) is False,
    )

    threads = get_adapter(Platform.THREADS)
    check(
        "threads needs only sessionid",
        await threads.has_live_session(FakeContext({"sessionid": "s"})) is True,
    )
    check(
        "threads without sessionid is not logged in",
        await threads.has_live_session(FakeContext({"ig_did": "x"})) is False,
    )

    check(
        "login/checkpoint paths are treated as interstitials",
        all(
            marker in facebook.interstitial_markers
            for marker in ("/login", "/checkpoint", "/two_step")
        ),
    )


def test_platform_access_flags() -> None:
    """requires_login must reflect measured behaviour, not guesswork."""
    print("\n[10] platform access flags")
    from app.core.exceptions import AuthenticationRequiredError, PlatformBlockedError
    from app.crawler import get_adapter

    tiktok = get_adapter(Platform.TIKTOK)
    check("tiktok does not require a login", tiktok.requires_login is False)
    check("tiktok asks for a visible window", tiktok.prefers_headed is True)

    for platform in (Platform.THREADS, Platform.INSTAGRAM, Platform.FACEBOOK):
        adapter = get_adapter(platform)
        check(f"{platform.value} still requires a login", adapter.requires_login is True)
        check(f"{platform.value} runs headless by default", adapter.prefers_headed is False)

    blocked = PlatformBlockedError("tiktok", "trả về trang lỗi", remedy="đặt SC_HEADLESS=false")
    check("blocked error keeps its remedy", blocked.remedy == "đặt SC_HEADLESS=false")
    check(
        "blocked is not an auth error",
        not isinstance(blocked, AuthenticationRequiredError),
    )
    check("BLOCKED status exists", TargetStatus.BLOCKED.value == "blocked")

    # Assert the property, not the tuned number: the budget now covers comment
    # XHRs as well as page loads, so the figure itself moves with tuning.
    check("tiktok sets its own request budget", tiktok.rate_per_minute is not None)
    check("tiktok gets no free burst", tiktok.rate_burst == 1)
    check(
        "other platforms use the global rate",
        get_adapter(Platform.THREADS).rate_per_minute is None,
    )

    from app.crawler.base import is_throttle_error

    throttled = Exception(
        "Page.goto: net::ERR_HTTP_RESPONSE_CODE_FAILURE at https://www.tiktok.com/x"
    )
    check("HTTP status failure is read as throttling", is_throttle_error(throttled) is True)
    check(
        "too-many-requests is read as throttling",
        is_throttle_error(Exception("net::ERR_TOO_MANY_REQUESTS")) is True,
    )
    check(
        "a timeout is NOT throttling",
        is_throttle_error(Exception("Timeout 45000ms exceeded")) is False,
    )


async def test_adaptive_cooldown() -> None:
    print("\n[11] adaptive rate-limit cooldown")
    from app.crawler.ratelimit import RateLimiter

    limiter = RateLimiter(per_minute=600, burst=5)
    check("not cooling down initially", limiter.cooling_down is False)

    first = limiter.penalise(10.0)
    check("first penalty lasts the base duration", first == 10.0, f"(got {first})")
    check("cooldown is now active", limiter.cooling_down is True)

    second = limiter.penalise(10.0)
    check("a repeat offence lasts longer", second == 20.0, f"(got {second})")

    limiter.relax()
    check("a success stops the escalation", limiter.penalise(10.0) == 10.0)

    # A cooldown must outrank available tokens.
    blocked = RateLimiter(per_minute=600, burst=5)
    blocked.penalise(0.6)
    started = time.monotonic()
    await blocked.acquire()
    waited = time.monotonic() - started
    check("acquire() waits out the cooldown", waited >= 0.5, f"(waited {waited:.2f}s)")

    free = RateLimiter(per_minute=600, burst=5)
    started = time.monotonic()
    await free.acquire()
    check(
        "no cooldown means no wait",
        time.monotonic() - started < 0.2,
    )


async def test_throttled_posts_get_a_second_pass() -> None:
    """A post turned away for going too fast must not be lost."""
    print("\n[12] throttled posts are retried, not dropped")
    from app.core.enums import TargetStatus  # noqa: F811
    from app.crawler import CrawlContext, CrawlOptions
    from app.crawler.ratelimit import RateLimiter
    from app.domain.entities import CrawlResult, PostRef
    from app.services.crawl_service import CrawlService

    throttle_error = Exception(
        "Could not open https://x/2: Page.goto: net::ERR_HTTP_RESPONSE_CODE_FAILURE at https://x/2"
    )
    hard_error = Exception("Could not open https://x/3: net::ERR_NAME_NOT_RESOLVED")

    class FlakyAdapter:
        platform = Platform.THREADS

        def __init__(self) -> None:
            self.calls: list[str] = []

        async def fetch_post(self, page, ref, ctx):  # noqa: ANN001 - test double
            self.calls.append(ref.external_id)
            if ref.external_id == "P2" and self.calls.count("P2") == 1:
                raise throttle_error          # succeeds on the second pass
            if ref.external_id == "P3":
                raise hard_error              # never recoverable
            return CrawlResult(post=_sample_post(5, 1), comments=[])

    async with session_scope() as session:
        job = await JobRepository(session).create_job(
            keyword="retry", platforms=[Platform.THREADS],
            search_filter=SearchFilter.TOP, max_posts=3,
            max_comments_per_post=5, include_comments=False,
        )
        job_id = job.id

    refs = [
        PostRef(platform=Platform.THREADS, external_id=f"P{n}", url=f"https://x/{n}", rank=n)
        for n in (1, 2, 3)
    ]
    adapter = FlakyAdapter()
    context = CrawlContext(
        job_id=job_id, platform=Platform.THREADS, keyword="retry",
        options=CrawlOptions(max_posts=3), cancel_event=asyncio.Event(),
        rate_limiter=RateLimiter(per_minute=6000, burst=50),
    )

    service = CrawlService(browser=None)  # the fake adapter never touches a browser
    await service._crawl_refs(job_id, Platform.THREADS, adapter, None, refs, context)

    check("throttled post was attempted twice", adapter.calls.count("P2") == 2,
          f"(calls={adapter.calls})")
    check("unrecoverable post was attempted once", adapter.calls.count("P3") == 1)

    async with session_scope() as session:
        target = (await JobRepository(session).list_targets(job_id))[0]
    check("throttled post counted as crawled, not failed", target.posts_crawled == 2,
          f"(crawled={target.posts_crawled})")
    check("only the genuine failure is counted", target.posts_failed == 1,
          f"(failed={target.posts_failed})")
    check("failure reason is recorded", "P3" in (target.message or ""),
          f"(message={target.message!r})")
    check("TargetStatus import still sane", TargetStatus.COMPLETED.value == "completed")


def test_export_filenames() -> None:
    """<keyword>_<platform>_<date>_<time> - readable at a glance."""
    print("\n[13] export file naming")
    import re
    from datetime import timezone as _tz
    from types import SimpleNamespace

    from app.db.repositories import PostFilters as _Filters
    from app.services.export_service import ExportService

    service = ExportService()
    post = lambda p: SimpleNamespace(platform=p)  # noqa: E731 - tiny stub

    name = service._build_filename(
        label="phê la mọng mơ", filters=_Filters(), posts=[post("tiktok")], fmt="xlsx"
    )
    check("diacritics folded to ASCII", name.startswith("phe-la-mong-mo_"), name)
    check("platform in the name", "_tiktok_" in name, name)
    check("extension kept", name.endswith(".xlsx"), name)
    check(
        "date and time appended",
        re.search(r"_\d{4}-\d{2}-\d{2}_\d{6}\.xlsx$", name) is not None,
        name,
    )

    mixed = service._build_filename(
        label="kw", filters=_Filters(), posts=[post("tiktok"), post("threads")], fmt="csv"
    )
    check("every platform present is listed", "_threads-tiktok_" in mixed, mixed)

    filtered = service._build_filename(
        label="kw", filters=_Filters(platform=Platform.FACEBOOK),
        posts=[post("tiktok")], fmt="json",
    )
    check("an explicit platform filter wins", "_facebook_" in filtered, filtered)

    empty = service._build_filename(label="", filters=_Filters(), posts=[], fmt="xlsx")
    check("no keyword falls back", empty.startswith("tat-ca_"), empty)
    check("no rows is marked", "_trong_" in empty, empty)

    # The name must read as local wall-clock time, not UTC.
    stamp = service._local_now()
    offset = stamp.utcoffset()
    check(
        "timestamp uses the configured timezone",
        offset is not None and offset != _tz.utc.utcoffset(None),
        f"(offset={offset})",
    )


def test_reply_expander_labels() -> None:
    """Facebook/Instagram reply-expander regex against real observed labels.

    Measured directly against a live Facebook post: Vietnamese uses two
    different words for "reply" depending on context ("cau tra loi" and
    "phan hoi"). A regex covering only one silently misses every thread
    using the other - which is exactly what the original TikTok-only
    pattern would have done here.
    """
    print("\n[14] Facebook/Instagram reply-expander label matching")
    from app.crawler.adapters.facebook import _REPLY_EXPANDER_RE as fb_re
    from app.crawler.adapters.instagram import _REPLY_EXPANDER_RE as ig_re

    real_labels = (
        "Xem 3 câu trả lời",
        "Xem 1 phản hồi",
        "View 5 replies",
    )
    for label in real_labels:
        check(f"facebook regex matches {label!r}", fb_re.search(label) is not None)
        check(f"instagram regex matches {label!r}", ig_re.search(label) is not None)

    non_matches = ("Viết bình luận...", "Thích", "Trả lời", "")
    for label in non_matches:
        check(
            f"facebook regex does not match {label!r}",
            fb_re.search(label) is None,
        )


async def test_has_state_checks_real_cookie_path() -> None:
    """Chromium keeps cookies at Default/Network/Cookies, not Default/Cookies.

    A profile with that file and nothing else used to be reported as
    "missing" in the UI, despite holding a real, working session.
    """
    print("\n[15] has_state() checks the real cookie path")
    from app.crawler import BrowserManager

    manager = BrowserManager()
    profile_dir = manager._settings.profile_dir(Platform.THREADS.value)
    network_dir = profile_dir / "Default" / "Network"
    network_dir.mkdir(parents=True, exist_ok=True)
    (network_dir / "Cookies").write_bytes(b"not a real sqlite file, just presence")

    check(
        "a profile with only Default/Network/Cookies is detected",
        manager.has_state(Platform.THREADS) is True,
    )


async def main() -> int:
    configure_logging("WARNING")
    test_number_parsing()
    test_metric_merge()
    job_id = await test_persistence()
    await test_job_status_rollup()
    await test_exports(job_id)
    test_json_body_parsing()
    test_threads_extraction()
    test_search_url_building()
    await test_session_cookie_rules()
    test_platform_access_flags()
    await test_adaptive_cooldown()
    await test_throttled_posts_get_a_second_pass()
    test_export_filenames()
    test_reply_expander_labels()
    await test_has_state_checks_real_cookie_path()
    await dispose_db()

    print(f"\n{'=' * 52}")
    if FAILED:
        print(f"{len(FAILED)} FAILED / {PASSED + len(FAILED)} checks")
        for label in FAILED:
            print(f"  - {label}")
        return 1
    print(f"All {PASSED} checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
